import pandas as pd
import tkinter as tk
from tkinter import messagebox
from tkinter import ttk
from datetime import datetime
from cronometro import CronometroApp
import openpyxl
from threading import Thread
import time

### UN WIDGET SE COMPONE DE 4 LINEAS DE CODIGO
### EL PASO CERO ES CREAR UN OBJETO QUE ALMACENE UN VALOR DE UN WIDGET, PERTENECIENTE A LA CLASE TKINTER, COMO StringVar()
### PRIMERO SE DECLARA EL WIDGET, INDICANDO EL TEXTO QUE CMOSTRARÁ Y EL RECIPIENTE QUE LO CONTENDRÁ
### Y LA SEGUNDA LINEA INDICA EN QUÉ POSICIÓN ESTARÁ EN EL WIDGET PADRE
### LA TERCERA, QUE AL IGUAL A LA LINEA CERO SON OPCIONALES DEPENDIENDO EL WIDGET

class Agregar_pedido:
    def __init__(self,root,tree,func_rtrn,pedidos):
        ### ABOUT
        ### ESTA ES UNA FUNCION CONSTRUCTOR DONDE SE DECLARAN LOS WIDGETS QUE FORMAN LA VENTANA
        
        self.pedidos = pedidos
        self.return_func = func_rtrn## Funcion de retorno a la ventana principal o ventana padre
        self.tree = tree# el treeview de la ventana principal o ventana padre
        ## CODIGO PARA ESTRAER FECHAY HORA
        ahora = datetime.now()
        self.fecha_y_hora = ahora.strftime("%Y-%m-%d %H:%M:%S")

        ## TODOS LOS WIDGETS A USAR ------------------------------------------------------------------------------------------------------------------------------------
        ## VENTANA
        self.root = tk.Toplevel(root,bg="#046546")# 0
        #### Primera Activity
        self.primer_activity = []
        ## FRAME
        self.primer_activity.append([ttk.LabelFrame(self.root, text="Detalles del Pedido", padding=(20, 10)),[2,0,3,3,"nsew"]])#0
        ## TITULO: HORA
        self.primer_activity.append([tk.Label(self.root,text=self.fecha_y_hora),[0,0,3,3,"nsew"]])
        ## FRAME: DIRECCIÓN DE ENTREGA
        self.direc_cliente_frame = ttk.LabelFrame(self.check_frame, text="Dirección de Entrega", padding=(20, 10))
        ## ENTRY: IRECCIÓN DE ENTREGA
        self.direc_cliente_entry = ttk.Entry(self.direc_cliente_frame)
        self.direc_cliente_entry.insert(0, "Purisima 55")
        self.direc_cliente_entry.bind("<Button-1>", lambda event: self.borrar_texto(self.direc_cliente_entry))# Asociar la función borrar_texto al evento clic
        ## NOMBRE DEL CLIENTE
        self.nombre_cliente_frame = ttk.LabelFrame(self.check_frame, text="Nombre Cliente", padding=(20, 10))
        self.nombre_cliente_entry = ttk.Entry(self.nombre_cliente_frame)
        self.nombre_cliente_entry.insert(0, "Tierno")
        self.nombre_cliente_entry.bind("<Button-1>", lambda event: self.borrar_texto(self.nombre_cliente_entry))# Asociar la función borrar_texto al evento clic
        ## NUMERO DEL CLIENTE
        self.num_cliente_frame = ttk.LabelFrame(self.check_frame, text="Numero celular", padding=(20, 10))
        self.num_cliente_entry = ttk.Entry(self.num_cliente_frame)
        self.num_cliente_entry.insert(0, "5510293483")
        self.num_cliente_entry.bind("<Button-1>", lambda event: self.borrar_texto(self.num_cliente_entry))# Asociar la función borrar_texto al evento clic
        ## TARIFA
        self.tarifa_frame = ttk.LabelFrame(self.check_frame, text="Tarifa $MXN", padding=(20, 10))
        self.tarifa_entry = ttk.Entry(self.tarifa_frame)
        self.tarifa_entry.insert(0, 10)
        self.tarifa_entry.bind("<Button-1>", lambda event: self.borrar_texto(self.tarifa_entry))# Asociar la función borrar_texto al evento clic
        ## Orden
        # Crear un widget Text
        self.orden_text_frame = ttk.LabelFrame(self.root, text="Productos")
        self.orden_text = tk.Text(self.orden_text_frame, height=10, width=40)
        # Create a Frame for the Radiobuttons
        self.radio_frame = ttk.LabelFrame(self.check_frame, text="Lista de articulos", padding=(20, 10))
        # Radiobuttons
        self.metodos_de_pago = ["Efectivo","Transferencia"]
        self.var_metodos_de_pago = tk.StringVar()
        self.var_metodos_de_pago.set(None)
        self.radio_metodos_de_pago = [tk.Radiobutton(self.radio_frame, text=metodo, variable=self.var_metodos_de_pago, value=metodo, command = lambda o = metodo:self.metodo_pago_func(o)) for i,metodo in enumerate(self.metodos_de_pago)]
        self.metodo_pago_label = tk.Label(self.radio_frame, text="No ha selecciónado\n metodo de pago")
        self.metodo_pago_frame = tk.Frame(self.radio_frame)
        self.paga_con = tk.StringVar()
        self.efectivo_widgets = [[tk.Entry(self.metodo_pago_frame,textvariable=self.paga_con),(0,0)]]
        self.transferencia_widgets = [[tk.Label(self.metodo_pago_frame, text= f"Numero de Tarjeta:\nXXXX-XXXX-XXXX-XXXX"),(0,0)],
                                 [tk.Button(self.metodo_pago_frame, text="Copiar Numero de cuenta"),(1,0)]]
        # BOTONES
        self.widgets_frame = ttk.Frame(self.check_frame)
        self.botones = [("Guardar",self.guardar_f),("Cancelar",self.cancelar_f)]
        self.botones_button = [ttk.Button(self.widgets_frame, text=i[0],command=i[1] ) for i in self.botones]
        #### SEgunda Activity
        

    def guardar_f(self):
        orden = str(self.orden_text.get("1.0", "end-1c"))
        print(f" Orden es : {orden}")
        if orden == "":
            messagebox.showwarning("Advertencia", "No has completado tu orden")
        else:
            dicc_orden = self.procesar_lista_compras(orden)
            print(dicc_orden)
            # "N° de Pedido","Fecha","Nombre del cliente","N° de Contacto","Recoleccion","Entrega","Producto","Paga con…","Costo del pedido","Tarifa","Cobro total","Dinero entregado a repartidor para compra","$ entregado al repartidor para cambio","$ total dado al repartidor","$ total recibido por el repartidor"]  

            direccion = self.direc_cliente_entry.get()
            nombre = self.nombre_cliente_entry.get()
            numero = self.num_cliente_entry.get()
            orden = self.orden_text.get("1.0", "end-1c")
            tarifa = self.tarifa_entry.get()
            metodo_pago = self.var_metodos_de_pago.get()
            paga_con = self.paga_con.get()
            list_cli = [self.fecha_y_hora,#0
                        nombre,#1
                        numero,#2
                        "No especificado",#3
                        direccion,#4
                        orden,#5
                        paga_con,#6
                        "No especificado",#7
                        tarifa,#8
                        "No especificado",#9
                        "No especificado",#10
                        metodo_pago # 11
                        ]

            dicc_cli = {"dict":["No especificado",str(self.fecha_y_hora),nombre,numero,"No especificado",direccion,orden,"No especificado","No especificado",tarifa,"No especificado","No especificado","No especificado","No especificado","No especificado"]}
            self.return_func(self.tree,list_cli)
    
        # Blocking widgets
        mallaopcion1.config(state=tk.DISABLED)
        chopcion1.config(state=tk.DISABLED)
        chopcion2.config(state=tk.DISABLED)
        chopcion3.config(state=tk.DISABLED)    
        chopcion4.config(state=tk.DISABLED)
        name_prod_entry.config(state=tk.DISABLED)
        no_prod_entry.config(state=tk.DISABLED)
        picto_prod_entry.config(state=tk.DISABLED)
        etiq_prod_entry.config(state=tk.DISABLED)
        mues_prod_entry.config(state=tk.DISABLED)
        caracmalla_entry.config(state=tk.DISABLED)
        self.root.destroy()
        
    def cancelar_f(self):
        self.root.destroy()

    ### METODO DE PAGO: RADIOBUTTONS COMMAND
    def metodo_pago_func(self,opcion):
        metodo = self.var_metodos_de_pago.get()
        self.metodo_pago_label.config(text=f"Pagará con:\n{metodo}")
        print(f"Metodo es : {metodo}")
        self.diccionarios_widgets = {"Efectivo":self.efectivo_widgets,"Transferencia":self.transferencia_widgets}
        for i,conjunto in enumerate(self.diccionarios_widgets.values()):
            print(f"conjunto {i} es : {conjunto}")
##            conjunto[0].grid_forget()
            for j, widget in enumerate(conjunto):
                print(f"widget {j} es: {widget}")
                widget[0].grid_forget()
        for i,lista in enumerate(self.diccionarios_widgets[opcion]):
            lista[0].grid(row=lista[1][0],column=lista[1][1])
    def procesar_lista_compras(self,cadena_lista):
        if any(char.isdigit() for char in cadena_lista):
            # La cadena contiene al menos un número
            print("La cadena contiene números. Realizar alguna acción aquí.")
            print(f"cadena : '{cadena_lista}'")
            lista_items = cadena_lista.split(', ')
            resultado = {}

            for item in lista_items:
                partes = item.split(' ')
                if len(partes) >= 2:
                    cantidad = int(partes[0])
                    producto = ' '.join(partes[1:])
                    resultado[producto] = cantidad
            return resultado
        else:
            # La cadena no contiene números, no hacer nada
            
            pass
        
    def draw(self):
        ## TODO LO QUE DIBUJA A LOS WIDGETS -------------------------------------------------------------------------------------------------------------------------------

        ## MAPA JERARQUICO DE WIDGETS
        ## root
        ##     fecha_y_hora
        ##     productos
        ##         orden: text
        ##     detalles del pedido/checkframe
        ##         direccion_de_entrega:entry
        ##         nombre:entry
        ##         numero:entry
        ##         metodo de pago: Radio_buttons
        
        self.check_frame.grid(row=2, column=0, padx=(3), pady=(3), sticky="nsew")
        self.hora_label.grid(row=0, column=0, padx=(3), pady=(3), sticky="nsew")

        self.direc_cliente_frame.grid(row=0, column=0, sticky="nsew")
        self.direc_cliente_entry.grid(row=0, column=0, padx=5, pady=(0, 10), sticky="ew")

        self.nombre_cliente_frame.grid(row=1, column=0, sticky="nsew")
        self.nombre_cliente_entry.grid(row=0, column=0, padx=5, pady=(0, 10), sticky="ew")

        self.num_cliente_frame.grid(row=1, column=1, sticky="nsew")
        self.num_cliente_entry.grid(row=0, column=0, padx=5, pady=(0, 10), sticky="ew")

        self.tarifa_frame.grid(row=0, column=1, sticky="nsew")
        self.tarifa_entry.grid(row=0, column=0, padx=5, pady=(0, 10), sticky="ew")
        
        self.orden_text_frame.grid(row=1, column=0, sticky="nsew")
        self.orden_text.pack(pady=10)

        self.radio_frame.grid(row=2, column=0, padx=(20, 10), pady=10, sticky="nsew")
        
        for i,radio in enumerate(self.radio_metodos_de_pago):            
            radio.grid(row=i, column=0, sticky="nsew")
        self.metodo_pago_label.grid(row=0,column=1)
        self.metodo_pago_frame.grid(row=1,column=1)
        for i,btn in enumerate(self.botones_button):
            btn.grid(row=0, column=i, sticky="nsew")
        self.widgets_frame.grid(row=3, column=0, sticky="nsew")

    def get_data(self,func):
        
        data = [self.fecha_y_hora,direccion, nombre, numero, orden]
        func(data)
        
    def borrar_texto(self,entry):
        # Borra el texto dentro del Entry al recibir un clic
        entry.delete(0, tk.END)
